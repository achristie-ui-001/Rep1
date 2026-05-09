import datetime
import os
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

from src.scoring.scorer import InflectionSignal

TIER_COLORS = {
    "Tier 1": "C8F7C5",
    "Tier 2": "FFF9C4",
    "Tier 3": "FFE0B2",
    "Filtered": "F5F5F5",
}

HEADER_FILL = PatternFill("solid", fgColor="1A237E")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

SIGNAL_COLUMNS = [
    ("Company", 28),
    ("Ticker", 9),
    ("Sector", 18),
    ("Date", 12),
    ("Quarter", 9),
    ("Inflection Type", 25),
    ("Tier", 9),
    ("Score", 8),
    ("Key Quote", 65),
    ("Source", 15),
    ("Novelty Score", 14),
    ("Dollar Qty?", 12),
    ("Analyst Q Count", 16),
    ("Speaker", 25),
    ("Speaker Role", 14),
    ("Transcript URL", 45),
    ("Next Action", 22),
]


def write_excel(signals: list[InflectionSignal], output_path: str, processing_log: list[dict] = None):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb = openpyxl.Workbook()

    _write_signals_sheet(wb, signals)
    _write_sector_summary(wb, signals)
    _write_processing_log(wb, processing_log or [])

    wb.save(output_path)
    return output_path


def _write_signals_sheet(wb, signals: list[InflectionSignal]):
    ws = wb.active
    ws.title = "Inflection Signals"
    ws.freeze_panes = "A2"

    # Header row
    for col_idx, (col_name, col_width) in enumerate(SIGNAL_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width

    ws.row_dimensions[1].height = 30

    # Sort by score descending
    sorted_signals = sorted(signals, key=lambda s: s.score, reverse=True)

    for row_idx, sig in enumerate(sorted_signals, start=2):
        fill = PatternFill("solid", fgColor=TIER_COLORS.get(sig.tier, "FFFFFF"))

        row_data = [
            sig.company_name,
            sig.ticker,
            sig.sector,
            sig.date,
            sig.quarter,
            sig.inflection_type,
            sig.tier,
            sig.score,
            sig.key_quote,
            sig.source,
            sig.novelty_score,
            "Yes" if sig.has_dollar_qty else "No",
            sig.analyst_q_count,
            sig.speaker_name,
            sig.speaker_role,
            sig.transcript_url,
            "",  # Next Action — blank for analyst
        ]

        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=(col_idx == 9))
            if col_idx == 16 and value:  # URL column
                cell.hyperlink = value
                cell.font = Font(color="1565C0", underline="single")

        ws.row_dimensions[row_idx].height = 40 if len(sig.key_quote) > 120 else 25

    # Score color scale (column H = 8)
    last_row = len(sorted_signals) + 1
    if last_row > 1:
        score_range = f"H2:H{last_row}"
        ws.conditional_formatting.add(
            score_range,
            ColorScaleRule(
                start_type="min", start_color="FF4444",
                mid_type="percentile", mid_value=50, mid_color="FFEE58",
                end_type="max", end_color="00C853",
            ),
        )

    # Auto-filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SIGNAL_COLUMNS))}1"


def _write_sector_summary(wb, signals: list[InflectionSignal]):
    ws = wb.create_sheet("Summary by Sector")
    headers = ["Sector", "Tier 1", "Tier 2", "Tier 3", "Total Signals", "Top Ticker", "Top Score"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = 18

    by_sector = defaultdict(list)
    for sig in signals:
        if sig.tier != "Filtered":
            by_sector[sig.sector].append(sig)

    row_idx = 2
    for sector, sector_signals in sorted(by_sector.items()):
        t1 = sum(1 for s in sector_signals if s.tier == "Tier 1")
        t2 = sum(1 for s in sector_signals if s.tier == "Tier 2")
        t3 = sum(1 for s in sector_signals if s.tier == "Tier 3")
        top_sig = max(sector_signals, key=lambda s: s.score)
        row = [sector, t1, t2, t3, len(sector_signals), top_sig.ticker, top_sig.score]
        for col_idx, val in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=val)
        row_idx += 1


def _write_processing_log(wb, log: list[dict]):
    ws = wb.create_sheet("Processing Log")
    headers = ["Ticker", "Company", "Source", "Status", "Transcripts Found", "Processing Time (s)", "Error Message"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        ws.column_dimensions[get_column_letter(col_idx)].width = 20

    for row_idx, entry in enumerate(log, start=2):
        row = [
            entry.get("ticker", ""),
            entry.get("company", ""),
            entry.get("source", ""),
            entry.get("status", ""),
            entry.get("transcripts_found", 0),
            round(entry.get("processing_time_s", 0), 1),
            entry.get("error_message", ""),
        ]
        for col_idx, val in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=val)


def make_output_path(template: str) -> str:
    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    return template.replace("{timestamp}", ts)
